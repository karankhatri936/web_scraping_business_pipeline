"""Tests for CSV and Excel report generation (temporary directories)."""

from __future__ import annotations

import csv

import pandas as pd
import pytest
from openpyxl import load_workbook

from analytics.analysis import compute_analysis
from data.cleaner import clean_products, records_to_dataframe
from reports.csv_report import write_products_csv
from reports.excel_report import write_excel_report


@pytest.fixture
def cleaned(sample_records):
    return clean_products(records_to_dataframe(sample_records)).frame


@pytest.fixture
def analysis(cleaned):
    return compute_analysis(cleaned, top_n=5)


def test_csv_created_and_contents(tmp_path, cleaned):
    path = write_products_csv(cleaned, tmp_path)
    assert path.exists() and path.suffix == ".csv"
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == len(cleaned)
    assert rows[0]["name"] == cleaned.iloc[0]["name"]
    assert "price" in rows[0] and "category" in rows[0]


def test_csv_output_directory_created(tmp_path, cleaned):
    nested = tmp_path / "deep" / "csv"
    path = write_products_csv(cleaned, nested)
    assert path.exists()


def test_csv_empty_dataset(tmp_path):
    path = write_products_csv(pd.DataFrame(), tmp_path, filename="empty.csv")
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader)
    assert "name" in header and "url" in header


def test_excel_created_with_expected_sheets(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    assert path.exists() and path.suffix == ".xlsx"
    wb = load_workbook(path)
    names = wb.sheetnames
    for expected in ("Dashboard", "Product Data", "Category Analysis",
                     "Top Products", "Report Metadata"):
        assert expected in names


def test_excel_dashboard_kpis(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    wb = load_workbook(path)
    ws = wb["Dashboard"]
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(4, 14)}
    assert labels["Total products"] == len(cleaned)
    assert labels["Average price"] is not None


def test_excel_product_data_rows(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    wb = load_workbook(path)
    ws = wb["Product Data"]
    header = [c.value for c in ws[1]]
    assert "name" in header and "price" in header
    data_rows = sum(1 for row in ws.iter_rows(min_row=2) if row[0].value)
    assert data_rows == len(cleaned)


def test_excel_category_analysis_table(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    wb = load_workbook(path)
    ws = wb["Category Analysis"]
    header = [c.value for c in ws[1]]
    assert "category" in header and "product_count" in header


def test_excel_top_products_sections(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    wb = load_workbook(path)
    ws = wb["Top Products"]
    values = {c.value for row in ws.iter_rows(max_col=1) for c in row}
    assert "Highest Rated Products" in values
    assert "Lowest Priced Products" in values


def test_excel_metadata_contains_source(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://books.toscrape.com/")
    wb = load_workbook(path)
    ws = wb["Report Metadata"]
    values = {c.value for row in ws.iter_rows(max_col=2) for c in row}
    assert "https://books.toscrape.com/" in values
    assert any(str(v).endswith("UTC") for v in values if isinstance(v, str))


def test_excel_empty_dataset_still_writes(tmp_path):
    empty = pd.DataFrame(columns=["name", "url", "price"])
    analysis = compute_analysis(empty)
    path = write_excel_report(
        analysis, empty, tmp_path, "https://example.com", filename="empty.xlsx"
    )
    wb = load_workbook(path)
    assert "Dashboard" in wb.sheetnames
    assert "Product Data" in wb.sheetnames


def test_excel_charts_present(tmp_path, cleaned, analysis):
    path = write_excel_report(analysis, cleaned, tmp_path, "https://example.com")
    wb = load_workbook(path)
    assert len(wb["Dashboard"]._charts) >= 2
