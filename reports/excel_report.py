"""Professional Excel reporting with OpenPyXL.

Generates a multi-sheet business report:

* ``Dashboard``         - KPI block + charts (category / availability / price)
* ``Product Data``      - full cleaned dataset as a filterable table
* ``Category Analysis`` - per-category aggregates as a table
* ``Top Products``      - highest rated + cheapest products
* ``Price Changes``     - movement vs the previous run (when history exists)
* ``Report Metadata``   - run facts, record counts, generation timestamp

Only data that exists in the dataset is written; no figures are invented.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from analytics.analysis import AnalysisResult
from utils.exceptions import ReportGenerationError
from utils.logger import get_logger

logger = get_logger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="1F4E78")
KPI_LABEL_FONT = Font(bold=True, color="404040")
KPI_VALUE_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRICE_FORMAT = "#,##0.00"
RATING_FORMAT = "0.0"
INTEGER_FORMAT = "#,##0"

SHEET_DASHBOARD = "Dashboard"
SHEET_PRODUCTS = "Product Data"
SHEET_CATEGORY = "Category Analysis"
SHEET_TOP = "Top Products"
SHEET_CHANGES = "Price Changes"
SHEET_META = "Report Metadata"


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _style_header(ws: Worksheet, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _autosize(ws: Worksheet, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)


def _table_name(sheet_title: str, start_row: int) -> str:
    """Build a valid, unique Excel table name for a sheet section.

    Excel imposes three constraints that a sheet title does not satisfy:
    table names may not contain spaces (or most punctuation), may not look
    like a cell reference, and must be unique within the workbook.
    """
    safe = re.sub(r"[^0-9A-Za-z_]", "", sheet_title)
    if not safe or not (safe[0].isalpha() or safe[0] == "_"):
        safe = f"Sheet{safe}"
    return f"Table{safe}R{start_row}"


def _write_table(ws: Worksheet, df: pd.DataFrame, start_row: int = 1) -> int:
    """Write a DataFrame as a styled Excel table; returns the next free row."""
    if df.empty:
        ws.cell(row=start_row, column=1, value="No data available for this section.")
        return start_row + 2
    for col_idx, column in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=col_idx, value=str(column))
    _style_header(ws, start_row, len(df.columns))
    for row_offset, record in enumerate(df.itertuples(index=False), start=1):
        for col_idx, value in enumerate(record, start=1):
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.isoformat()
            cell = ws.cell(row=start_row + row_offset, column=col_idx, value=value)
            cell.border = BORDER
    table = Table(
        displayName=_table_name(ws.title, start_row),
        ref=f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    ws.add_table(table)
    return start_row + len(df) + 2


def _add_bar_chart(
    ws: Worksheet, title: str, data_ws: Worksheet, min_row: int, max_row: int,
    col: int, anchor: str, cats_col: int = 1, y_title: str = "Products",
) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.title = y_title
    chart.legend = None
    chart.height = 7.5
    chart.width = 15
    data = Reference(data_ws, min_col=col, min_row=min_row, max_row=max_row)
    cats = Reference(data_ws, min_col=cats_col, min_row=min_row + 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def _add_pie_chart(
    ws: Worksheet, title: str, data_ws: Worksheet, min_row: int, max_row: int,
    anchor: str,
) -> None:
    chart = PieChart()
    chart.title = title
    chart.height = 7.5
    chart.width = 12
    data = Reference(data_ws, min_col=2, min_row=min_row, max_row=max_row)
    labels = Reference(data_ws, min_col=1, min_row=min_row + 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    ws.add_chart(chart, anchor)


def _write_dashboard(wb: Workbook, analysis: AnalysisResult) -> None:
    ws = wb.active
    ws.title = SHEET_DASHBOARD
    kpis = analysis.kpis
    ws.cell(row=1, column=1, value="Product Monitoring Dashboard").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Generated {_timestamp()}").font = KPI_LABEL_FONT

    labels_values: list[tuple[str, Any, str | None]] = [
        ("Total products", kpis.get("total_products", 0), INTEGER_FORMAT),
        ("Unique products", kpis.get("unique_products", 0), INTEGER_FORMAT),
        ("Categories", kpis.get("categories", 0), INTEGER_FORMAT),
        ("In stock", kpis.get("in_stock", 0), INTEGER_FORMAT),
        ("Out of stock", kpis.get("out_of_stock", 0), INTEGER_FORMAT),
        ("Average price", kpis.get("avg_price"), PRICE_FORMAT),
        ("Median price", kpis.get("median_price"), PRICE_FORMAT),
        ("Min price", kpis.get("min_price"), PRICE_FORMAT),
        ("Max price", kpis.get("max_price"), PRICE_FORMAT),
        ("Average rating", kpis.get("avg_rating"), RATING_FORMAT),
    ]
    row = 4
    for label, value, number_format in labels_values:
        ws.cell(row=row, column=1, value=label).font = KPI_LABEL_FONT
        value_cell = ws.cell(
            row=row, column=2, value=value if value is not None else "n/a"
        )
        value_cell.font = KPI_VALUE_FONT
        if value is not None and number_format:
            value_cell.number_format = number_format
        row += 1

    charts_start = row + 2
    if not analysis.by_category.empty:
        cat_sheet = wb.create_sheet(SHEET_CATEGORY)
        _write_table(cat_sheet, analysis.by_category)
        _autosize(cat_sheet)
        max_row = min(1 + len(analysis.by_category), 11)  # top 10 categories
        _add_bar_chart(
            ws, "Products by Category (Top 10)", cat_sheet, 1, max_row, 2,
            f"A{charts_start}", y_title="Products",
        )
        _add_bar_chart(
            ws, "Average Price by Category (Top 10)", cat_sheet, 1, max_row, 5,
            f"J{charts_start}", y_title="Avg price",
        )
    if not analysis.availability.empty:
        avail_sheet = wb.create_sheet("_chartdata_availability")
        _write_table(avail_sheet, analysis.availability)
        _add_pie_chart(
            ws, "Availability Breakdown", avail_sheet, 1,
            1 + len(analysis.availability), f"S{charts_start}",
        )
        avail_sheet.sheet_state = "hidden"
    if not analysis.price_distribution.empty:
        dist_sheet = wb.create_sheet("_chartdata_price")
        _write_table(dist_sheet, analysis.price_distribution)
        _add_bar_chart(
            ws, "Price Distribution", dist_sheet, 1,
            1 + len(analysis.price_distribution), 2, f"A{charts_start + 16}",
            y_title="Products",
        )
        dist_sheet.sheet_state = "hidden"
    ws.sheet_view.showGridLines = False


def _write_metadata(
    wb: Workbook,
    analysis: AnalysisResult,
    source_url: str,
    extra: dict[str, Any] | None = None,
) -> None:
    ws = wb.create_sheet(SHEET_META)
    ws.cell(row=1, column=1, value="Report Metadata").font = TITLE_FONT
    rows: list[tuple[str, Any]] = [
        ("Generated at (UTC)", _timestamp()),
        ("Data source", source_url),
        ("Total rows in dataset", analysis.kpis.get("total_products", 0)),
        ("Report includes price data", analysis.kpis.get("avg_price") is not None),
        ("Report includes rating data", analysis.kpis.get("avg_rating") is not None),
    ]
    if extra:
        rows.extend(sorted(extra.items()))
    for idx, (label, value) in enumerate(rows, start=3):
        ws.cell(row=idx, column=1, value=label).font = KPI_LABEL_FONT
        ws.cell(row=idx, column=2, value=value)
    _autosize(ws)


def write_excel_report(
    analysis: AnalysisResult,
    products: pd.DataFrame,
    output_dir: Path | str,
    source_url: str,
    extra_metadata: dict[str, Any] | None = None,
    filename: str | None = None,
) -> Path:
    """Build the full multi-sheet Excel report; returns the written path."""
    output_dir = Path(output_dir)
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        _write_dashboard(wb, analysis)

        product_sheet = wb.create_sheet(SHEET_PRODUCTS)
        columns = [
            "name", "price", "currency", "category", "availability",
            "availability_status", "rating", "rating_band", "price_bucket",
            "url", "image_url", "source_product_id", "scraped_at",
        ]
        if not products.empty:
            view = products[[c for c in columns if c in products.columns]]
        else:
            view = pd.DataFrame(columns=columns)
        _write_table(product_sheet, view)
        product_sheet.freeze_panes = "A2"
        _autosize(product_sheet)

        top_sheet = wb.create_sheet(SHEET_TOP)
        top_sheet.cell(row=1, column=1, value="Highest Rated Products").font = TITLE_FONT
        next_row = _write_table(top_sheet, analysis.top_rated, start_row=2) + 1
        top_sheet.cell(row=next_row, column=1, value="Lowest Priced Products").font = TITLE_FONT
        _write_table(top_sheet, analysis.cheapest, start_row=next_row + 1)
        _autosize(top_sheet)

        if not analysis.price_changes.empty:
            changes_sheet = wb.create_sheet(SHEET_CHANGES)
            _write_table(changes_sheet, analysis.price_changes)
            _autosize(changes_sheet)

        _write_metadata(wb, analysis, source_url, extra_metadata)

        path = output_dir / (
            filename
            or f"product_report_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        wb.save(path)
        logger.info("Excel report written: %s", path)
        return path
    except (OSError, ValueError) as exc:
        raise ReportGenerationError(f"Could not write Excel report: {exc}") from exc
